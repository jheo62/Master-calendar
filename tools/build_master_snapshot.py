#!/usr/bin/env python3

import argparse
import json
from datetime import date, datetime, time, timezone
from pathlib import Path

import openpyxl


FIELDS = [
    "ID_Partido",
    "Temporada",
    "Visitante",
    "Local",
    "Arena",
    "Ciudad",
    "Estado_Provincia",
    "País",
    "Fase",
    "TV_Normalizada",
    "Streaming_Normalizado",
    "Grupo_NBA_Cup",
    "Evento_Especial",
    "Estado",
    "Es_NBA_Cup",
    "Es_Partido_Especial",
    "Zona_Horaria_Arena",
    "URL_NBA",
    "Fecha_Hora_UTC",
]


def clean(value):
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def yes_no(value):
    if isinstance(value, bool):
        return value

    if value is None:
        return False

    text = str(value).strip().lower()

    return (
        text == "sí"
        or text == "si"
        or text.startswith("sí ")
        or text.startswith("si ")
        or text.startswith("sí—")
        or text.startswith("si—")
        or text.startswith("sí —")
        or text.startswith("si —")
    )


def utc_iso(value):
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        else:
            value = value.astimezone(timezone.utc)

        return value.strftime("%Y-%m-%dT%H:%M:%SZ")

    if isinstance(value, date):
        dt = datetime.combine(value, time.min, tzinfo=timezone.utc)
        return dt.strftime("%Y-%m-%dT%H:%M:%SZ")

    if value is None:
        return None

    text = str(value).strip()

    if text.endswith("Z"):
        return text

    try:
        dt = datetime.fromisoformat(text)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        else:
            dt = dt.astimezone(timezone.utc)

        return dt.strftime("%Y-%m-%dT%H:%M:%SZ")
    except ValueError:
        raise ValueError(f"Fecha_Hora_UTC no reconocida: {value!r}")


def main():
    parser = argparse.ArgumentParser(
        description="Construye master_snapshot JSON desde Base Maestra XLSX."
    )
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    xlsx = Path(args.xlsx)
    output = Path(args.output)

    if not xlsx.exists():
        raise SystemExit(f"ERROR: no existe XLSX: {xlsx}")

    wb = openpyxl.load_workbook(
        xlsx,
        read_only=True,
        data_only=True,
    )

    candidates = []

    for ws in wb.worksheets:
        first = next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True),
            ()
        )

        headers = [
            str(v).strip() if v is not None else ""
            for v in first
        ]

        if "ID_Partido" in headers:
            candidates.append(ws)

    if not candidates:
        raise SystemExit("ERROR: no existe hoja con ID_Partido.")

    ws = max(candidates, key=lambda w: w.max_row)

    headers = [
        str(v).strip() if v is not None else ""
        for v in next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True)
        )
    ]

    missing = [f for f in FIELDS if f not in headers]

    if missing:
        raise SystemExit(
            "ERROR: faltan columnas requeridas: "
            + ", ".join(missing)
        )

    events = []

    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in values):
            continue

        row = dict(zip(headers, values))

        event = {
            "ID_Partido": clean(row["ID_Partido"]),
            "Temporada": clean(row["Temporada"]),
            "Visitante": clean(row["Visitante"]),
            "Local": clean(row["Local"]),
            "Arena": clean(row["Arena"]),
            "Ciudad": clean(row["Ciudad"]),
            "Estado_Provincia": clean(row["Estado_Provincia"]),
            "País": clean(row["País"]),
            "Fase": clean(row["Fase"]),
            "TV_Normalizada": clean(row["TV_Normalizada"]),
            "Streaming_Normalizado": clean(
                row["Streaming_Normalizado"]
            ),
            "Grupo_NBA_Cup": clean(row["Grupo_NBA_Cup"]),
            "Evento_Especial": clean(row["Evento_Especial"]),
            "Estado": clean(row["Estado"]),
            "Es_NBA_Cup": yes_no(row["Es_NBA_Cup"]),
            "Es_Partido_Especial": yes_no(
                row["Es_Partido_Especial"]
            ),
            "Zona_Horaria_Arena": clean(
                row["Zona_Horaria_Arena"]
            ),
            "URL_NBA": clean(row["URL_NBA"]),
            "Fecha_Hora_UTC": utc_iso(
                row["Fecha_Hora_UTC"]
            ),
        }

        events.append(event)

    wb.close()

    ids = [e["ID_Partido"] for e in events]

    if len(events) != 1200:
        raise SystemExit(
            f"ERROR: se esperaban 1200 eventos; encontrados {len(events)}"
        )

    if len(ids) != len(set(ids)):
        raise SystemExit("ERROR: existen ID_Partido duplicados.")

    payload = {
        "season": "2026-27",
        "generated_from": xlsx.name,
        "events": events,
    }

    output.parent.mkdir(parents=True, exist_ok=True)

    output.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2
        ) + "\n",
        encoding="utf-8",
    )

    print("SNAPSHOT CONSTRUIDO: OK")
    print("Hoja:", ws.title)
    print("Eventos:", len(events))
    print("IDs únicos:", len(set(ids)))
    print("Salida:", output)


if __name__ == "__main__":
    main()
